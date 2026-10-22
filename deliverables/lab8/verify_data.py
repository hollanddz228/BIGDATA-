"""Optional data audit: Python 3 + openpyxl. No Power BI execution is claimed.

Run from any directory: python path/to/verify_data.py
Writes DATA_CHECKS.json next to this script using the supplied source workbooks.
"""

import csv
import hashlib
import json
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent


def read_sheets(path, names):
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = {}
    try:
        for name in names:
            it = workbook[name].iter_rows(values_only=True)
            headers = next(it)
            assert len(set(headers)) == len(headers)
            result[name] = [dict(zip(headers, row)) for row in it if any(v is not None for v in row)]
    finally:
        workbook.close()
    return result


def index(rows, key):
    result = {r[key]: r for r in rows}
    assert len(result) == len(rows), f"Duplicate key: {key}"
    assert None not in result, f"Blank key: {key}"
    return result


def equivalent(a, b):
    if str(a) == str(b):
        return True
    try:
        return Decimal(str(a)) == Decimal(str(b))
    except InvalidOperation:
        return False


def compare_csv(path, excel_rows):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    assert len(csv_rows) == len(excel_rows)
    for csv_row, excel_row in zip(csv_rows, excel_rows):
        assert set(csv_row) == set(excel_row)
        for key, value in csv_row.items():
            assert equivalent(value, excel_row[key]), (path.name, key, value, excel_row[key])


sales = read_sheets(ROOT / "part1" / "Construct a creature.xlsx", ["Purchase", "Centre", "Town", "Region"])
purchase = sales["Purchase"]
index(purchase, "PurchaseId")
centres = index(sales["Centre"], "CentreId")
towns = index(sales["Town"], "TownId")
regions = index(sales["Region"], "RegionId")
assert all(row["TownId"] in towns for row in centres.values())
assert all(row["RegionId"] in regions for row in towns.values())
romford = Counter()
region_totals = Counter()
for row in purchase:
    centre = centres[row["CentreId"]]
    town = towns[centre["TownId"]]
    region = regions[town["RegionId"]]
    assert isinstance(row["Quantity"], (int, float))
    region_totals[region["RegionName"]] += row["Quantity"]
    if town["TownName"] == "Romford":
        romford[centre["CentreName"]] += row["Quantity"]
assert dict(romford) == {"The Brewery": 403, "The Liberty Shopping Centre": 343, "The Mall Romford": 119}
assert sum(row["Quantity"] for row in purchase) == 74415

voting = read_sheets(ROOT / "part2" / "Voting.xlsx", ["Voting"])["Voting"]
country_rows = read_sheets(ROOT / "part2" / "Regions.xlsx", ["Regions"])["Regions"]
compare_csv(ROOT / "part2" / "Voting.csv", voting)
compare_csv(ROOT / "part2" / "Regions.csv", country_rows)
area_index = index(voting, "Area_Code")
country_index = index(country_rows, "Region")
countries = defaultdict(lambda: {"Leave": 0, "Remain": 0, "Valid_Votes": 0, "areas": 0})
for row in voting:
    assert row["Leave"] + row["Remain"] == row["Valid_Votes"] > 0
    assert row["Region"] in country_index
    country = "Gibraltar" if row["Area_Code"] == "GI" else country_index[row["Region"]]["Country"]
    countries[country]["areas"] += 1
    for col in ["Leave", "Remain", "Valid_Votes"]:
        countries[country][col] += row[col]
for values in countries.values():
    values["LeaveShare"] = values["Leave"] / values["Valid_Votes"]
top10 = sorted(voting, key=lambda r: r["Leave"] / r["Valid_Votes"], reverse=True)[:10]
totals = {col: sum(row[col] for row in voting) for col in ["Leave", "Remain", "Valid_Votes"]}
assert len(voting) == 382 and len(country_rows) == 12
assert totals == {"Leave": 17410742, "Remain": 16141241, "Valid_Votes": 33551983}
assert top10[0]["Area"] == "Boston"
assert area_index["GI"]["Region"] == "South West"

report = {
    "validation_scope": "Source data and independent computations; Power BI Desktop and DAX were not executed.",
    "sales_rows": {name: len(rows) for name, rows in sales.items()},
    "sales_keys_unique_and_foreign_keys_matched": True,
    "total_quantity": sum(row["Quantity"] for row in purchase),
    "london_and_south_east_quantity": region_totals["London"] + region_totals["South East"],
    "romford": dict(romford),
    "romford_total": sum(romford.values()),
    "voting_rows": len(voting),
    "regions_rows": len(country_rows),
    "excel_csv_equivalent": True,
    "area_codes_unique": True,
    "all_voting_regions_mapped": True,
    "valid_votes_equal_leave_plus_remain_for_every_area": True,
    "totals": totals,
    "national_leave_share": totals["Leave"] / totals["Valid_Votes"],
    "top10": [{"Area": row["Area"], "PercentLeave": row["Leave"] / row["Valid_Votes"]} for row in top10],
    "countries_gibraltar_separate": dict(countries),
    "source_file_sha256": {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for part in ["part1", "part2"] for path in sorted((ROOT / part).iterdir()) if path.is_file()},
}
(ROOT / "DATA_CHECKS.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"sales_rows": report["sales_rows"], "romford": report["romford_total"], "voting_rows": len(voting), "excel_csv_equivalent": True, "status": "passed"}, ensure_ascii=False))
